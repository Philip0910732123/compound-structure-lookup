#!/usr/bin/env python3
"""API 冒烟测试：针对运行中的 exe (127.0.0.1:5173) 测试 Excel 上传 + 查询"""
import requests
import json
import os
import time

BASE = "http://127.0.0.1:5173"

def test_homepage():
    """测试首页可访问"""
    r = requests.get(BASE + "/", timeout=5)
    assert r.status_code == 200
    assert "化合物结构信息查询工具" in r.text
    assert "上传 Excel" in r.text
    assert "结构式" in r.text
    print("[PASS] test_homepage - 页面包含 '上传 Excel' 和 '结构式'")

def test_excel_upload():
    """测试 Excel 上传接口"""
    from openpyxl import Workbook
    
    # 创建测试 Excel：数据在 C 列（验证自动检测列）
    test_file = "api_test_input.xlsx"
    wb = Workbook()
    ws = wb.active
    ws['A1'] = '序号'
    ws['A2'] = 1
    ws['A3'] = 2
    ws['B1'] = '备注'
    ws['B2'] = '无关数据'
    ws['C1'] = '化合物名称'
    ws['C2'] = 'aspirin'
    ws['C3'] = 'caffeine'
    wb.save(test_file)
    
    # 上传
    with open(test_file, "rb") as f:
        r = requests.post(BASE + "/api/upload", files={"file": f}, timeout=10)
    
    os.remove(test_file)
    
    assert r.status_code == 200, f"上传失败: {r.status_code} {r.text}"
    data = r.json()
    compounds = data.get("compounds", [])
    print(f"  解析到 {len(compounds)} 个化合物: {compounds}")
    assert len(compounds) == 2, f"期望 2 个，实际 {len(compounds)}"
    assert compounds[0] == 'aspirin'
    assert compounds[1] == 'caffeine'
    print("[PASS] test_excel_upload - C列数据自动检测，跳过表头")

def test_query_with_image():
    """测试查询接口返回结构图片 URL"""
    r = requests.post(
        BASE + "/api/query",
        json={"compounds": "aspirin"},
        stream=True,
        timeout=60,
    )
    assert r.status_code == 200, f"查询失败: {r.status_code}"
    
    # 读取 SSE 流
    buffer = ""
    found_row = None
    for chunk in r.iter_content(chunk_size=1024, decode_unicode=True):
        buffer += chunk
        while "\n" in buffer:
            line, buffer = buffer.split("\n", 1)
            line = line.strip()
            if line.startswith("data: "):
                data = json.loads(line[6:])
                if data.get("phase") == "done" and data.get("row"):
                    found_row = data["row"]
            elif line.startswith("event: done"):
                break
    
    assert found_row is not None, "未收到查询结果"
    print(f"  SMILES: {found_row.get('SMILES')}")
    print(f"  CAS: {found_row.get('CAS号')}")
    print(f"  结构图片: {found_row.get('结构图片')}")
    assert found_row.get("结构图片"), "结构图片 URL 不应为空"
    assert "pubchem" in found_row["结构图片"] or "cactus" in found_row["结构图片"]
    print("[PASS] test_query_with_image - 查询结果包含结构图片 URL")

if __name__ == "__main__":
    print("=== exe API 冒烟测试 (127.0.0.1:5173) ===")
    print()
    
    test_homepage()
    print()
    
    test_excel_upload()
    print()
    
    test_query_with_image()
    print()
    
    print("=== 全部 API 测试通过 ===")
