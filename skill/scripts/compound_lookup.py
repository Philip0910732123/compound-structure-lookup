#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
化合物结构信息查询脚本 v3.3-skill

v3.3-skill 变更（相比 v3.2，同步 exe v3.5.7 改进）：
  R. Excel 输入解析 — parse_excel() 智能检测化合物名称列
  S. 结构图片 URL — PubChem CID 优先 / CACTUS SMILES 兜底
  T. 进度回调机制 — run(progress_callback=...) 供 AI 平台实时展示
  U. CSV/JSON 导出 — 无 pandas 依赖，csv 模块轻量导出
  V. CAS 预计算 — 富集前预判状态，SSE 推送时状态不为空
  W. parse_input — 逗号/换行分隔自动解析

数据获取优先级:
  Level 0:  RDKit 本地计算 InChIKey + InChI（不消耗任何外部 API）
  Level 1:  NCI CACTUS（SMILES/InChIKey → CAS + SMILES + IUPAC + 分子式 + 分子量 + InChIKey）
  Level 2:  PubChem PUG REST + PUG-View + DepositorRegistryID + annotations
  Level 2b: NIST Chemistry WebBook
  Level 2c: Wikidata SPARQL 批量查询
  Level 2d: EPA CompTox Dashboard（兜底）
"""

import sys
import os
import json
import time
import re
import csv
import io
import requests
from urllib.parse import quote
from pathlib import Path

OUT_DIR = os.environ.get("EUREKA_PYTHON_OUTPUT_DIR", ".")
TIMEOUT = 12
HEADERS = {"User-Agent": "CompoundLookup/3.3-skill"}

# ─────────────────────────────────────────────
# Level 0: RDKit 本地 InChIKey 计算
# ─────────────────────────────────────────────

_RDKIT_AVAILABLE = False
try:
    from rdkit import Chem
    _RDKIT_AVAILABLE = True
except ImportError:
    pass


def compute_inchikey_local(smiles):
    """用 RDKit 本地计算 InChIKey 和 InChI，不消耗任何外部 API"""
    if not _RDKIT_AVAILABLE or not smiles:
        return None, None
    try:
        mol = Chem.MolFromSmiles(smiles)
        if mol is None:
            return None, None
        inchikey = Chem.MolToInchiKey(mol)
        inchi = Chem.MolToInchi(mol)
        return inchikey, inchi
    except Exception:
        return None, None

# ─────────────────────────────────────────────
# 工具函数
# ─────────────────────────────────────────────

def is_smiles(text: str) -> bool:
    smiles_chars = set("CNOPSFClBrIHcnopsb=#@+\\/-[]()0123456789.%")
    if len(text) < 3:
        return False
    char_ratio = sum(1 for c in text if c in smiles_chars) / len(text)
    has_organic = bool(re.search(r'[CNOPSFcnops]', text))
    has_bond = bool(re.search(r'[=#\-\\/@\[\]().0-9]', text))
    return char_ratio > 0.85 and has_organic and has_bond


def safe_get(url, params=None, retries=3, timeout=None):
    """带指数退避的重试，处理 PubChem 503 等临时错误"""
    _timeout = timeout or TIMEOUT
    for attempt in range(retries):
        try:
            r = requests.get(url, params=params, headers=HEADERS, timeout=_timeout)
            if r.status_code == 200:
                return r
            if r.status_code == 503 and attempt < retries - 1:
                wait = 2 ** attempt
                time.sleep(wait)
                continue
            return r
        except Exception:
            if attempt < retries - 1:
                time.sleep(0.5 * (attempt + 1))
    return None


def validate_cas(cas_str):
    """验证 CAS 号校验码"""
    if not cas_str:
        return False
    m = re.match(r'^(\d{2,7})-(\d{2})-(\d)$', str(cas_str).strip())
    if not m:
        return False
    digits = m.group(1) + m.group(2)
    check = int(m.group(3))
    total = sum(int(d) * (i + 1) for i, d in enumerate(reversed(digits)))
    return total % 10 == check


def clean_cas(text):
    """从文本中提取标准 CAS 号并优先返回校验码通过的"""
    if not text:
        return None
    for line in text.strip().splitlines():
        line = line.strip()
        if re.match(r'^\d{2,7}-\d{2}-\d$', line):
            if validate_cas(line):
                return line
    for m in re.finditer(r'\b(\d{2,7}-\d{2}-\d)\b', text):
        cas = m.group(1)
        if validate_cas(cas):
            return cas
    m = re.search(r'\b(\d{2,7}-\d{2}-\d)\b', text)
    return m.group(1) if m else None


# ─────────────────────────────────────────────
# 名称翻译
# ─────────────────────────────────────────────

DICT_MAP = {
    "阿司匹林": "aspirin", "布洛芬": "ibuprofen",
    "对乙酰氨基酚": "acetaminophen", "青霉素": "penicillin",
    "咖啡因": "caffeine", "乙醇": "ethanol",
    "葡萄糖": "glucose", "胆固醇": "cholesterol",
    "维生素C": "vitamin C", "多巴胺": "dopamine",
    "肾上腺素": "epinephrine", "吗啡": "morphine",
    "氯化钠": "sodium chloride", "硫酸": "sulfuric acid",
    "盐酸": "hydrochloric acid", "奥希替尼": "osimertinib",
    "伏美替尼": "furmonertinib", "达沙替尼": "dasatinib",
    "伊马替尼": "imatinib", "吉非替尼": "gefitinib",
    "厄洛替尼": "erlotinib", "索拉非尼": "sorafenib",
    "拉帕替尼": "lapatinib", "克唑替尼": "crizotinib",
}


def translate_name(name):
    for zh, en in DICT_MAP.items():
        if name == zh:
            return en
        if zh in name:
            return name.replace(zh, en)
    if not re.search(r'[\u4e00-\u9fff]', name):
        return name
    r = safe_get(
        "https://api.mymemory.translated.net/get",
        params={"q": name, "langpair": "zh|en"}
    )
    if r:
        try:
            data = r.json()
            translated = data.get("responseData", {}).get("translatedText", "")
            if translated and translated.lower() != name.lower():
                return translated
        except Exception:
            pass
    return name

# ─────────────────────────────────────────────
# Level 1: NCI CACTUS (SMILES + InChIKey)
# ─────────────────────────────────────────────

CACTUS_BASE = "https://cactus.nci.nih.gov/chemical/structure"


def cactus_get(identifier, rep):
    url = f"{CACTUS_BASE}/{requests.utils.quote(identifier, safe='')}/{rep}"
    r = safe_get(url)
    if r:
        text = r.text.strip()
        if text and "Page not found" not in text and "error" not in text.lower():
            return text
    return None


def fetch_from_cactus(name, en_name, smiles_hint, inchikey_hint=None):
    result = {"smiles": None, "cas": None, "iupac": None,
              "formula": None, "mw": None, "inchikey": None, "source": None}

    identifiers = []
    if smiles_hint:
        identifiers.append((smiles_hint, "smiles"))
    if inchikey_hint:
        identifiers.append((inchikey_hint, "inchikey"))
    if en_name and en_name != name:
        identifiers.append((en_name, "en_name"))
    identifiers.append((name, "raw_name"))

    for ident, label in identifiers:
        smiles = cactus_get(ident, "smiles")
        if smiles:
            result["smiles"] = smiles.splitlines()[0].strip()
            result["source"] = f"CACTUS({label})"
            ref = result["smiles"]
            cas_raw = cactus_get(ref, "cas")
            result["cas"] = clean_cas(cas_raw)
            iupac = cactus_get(ref, "iupac_name")
            result["iupac"] = iupac.splitlines()[0].strip() if iupac else None
            formula = cactus_get(ref, "formula")
            result["formula"] = formula.strip() if formula else None
            mw = cactus_get(ref, "mw")
            result["mw"] = mw.strip() if mw else None
            inchikey = cactus_get(ref, "stdinchikey")
            result["inchikey"] = inchikey.strip().replace("InChIKey=", "") if inchikey else None
            break

    if not result["cas"]:
        for ident, label in identifiers:
            cas_raw = cactus_get(ident, "cas")
            cas = clean_cas(cas_raw)
            if cas:
                result["cas"] = cas
                if not result["source"]:
                    result["source"] = f"CACTUS-cas({label})"
                break

    if not result["cas"] and inchikey_hint:
        cas_raw = cactus_get(inchikey_hint, "cas")
        cas = clean_cas(cas_raw)
        if cas:
            result["cas"] = cas
            if not result["source"]:
                result["source"] = "CACTUS-inchikey-cas"

    return result


# ─────────────────────────────────────────────
# Level 2: PubChem (PUG REST + PUG-View + DepositorRegistryID + annotations)
# ─────────────────────────────────────────────

PUBCHEM_BASE = "https://pubchem.ncbi.nlm.nih.gov/rest/pug"
PUGVIEW_BASE = "https://pubchem.ncbi.nlm.nih.gov/rest/pug_view"


def fetch_cas_from_pug_view(cid):
    r = safe_get(f"{PUGVIEW_BASE}/data/compound/{cid}/JSON", timeout=20)
    if not r:
        return None
    try:
        data = r.json()
    except Exception:
        return None

    cas_candidates = []

    def walk_sections(sections):
        for sec in sections:
            if "data" in sec:
                for info in sec["data"]:
                    val = info.get("value", {})
                    sval = val.get("string_with_markup", [])
                    if isinstance(sval, list):
                        for sm in sval:
                            s = sm.get("string", "") if isinstance(sm, dict) else str(sm)
                            cas = clean_cas(s)
                            if cas:
                                cas_candidates.append(cas)
                    num = val.get("number", "")
                    if num:
                        cas = clean_cas(str(num))
                        if cas:
                            cas_candidates.append(cas)
            if "section" in sec:
                walk_sections(sec["section"])

    sections = data.get("Section", {}).get("TOC", [])
    if sections:
        walk_sections(sections)

    for info in data.get("Compound", {}).get("Information", []):
        val = info.get("Value", {})
        sval = val.get("StringWithMarkup", [])
        if isinstance(sval, list):
            for sm in sval:
                s = sm.get("String", "") if isinstance(sm, dict) else str(sm)
                cas = clean_cas(s)
                if cas:
                    cas_candidates.append(cas)

    for cas in cas_candidates:
        if validate_cas(cas):
            return cas
    return cas_candidates[0] if cas_candidates else None


def fetch_cas_from_depositor_xref(cid):
    r = safe_get(f"{PUBCHEM_BASE}/compound/cid/{cid}/xrefs/DepositorRegistryID/JSON")
    if not r:
        return None
    try:
        ids = r.json()["InformationList"]["Information"][0]["DepositorRegistryID"]
        for rid in ids:
            if re.match(r'^\d{2,7}-\d{2}-\d$', rid):
                if validate_cas(rid):
                    return rid
        for rid in ids:
            if re.match(r'^\d{2,7}-\d{2}-\d$', rid):
                return rid
    except Exception:
        pass
    return None


def fetch_cas_from_annotations(cid):
    r = safe_get(f"{PUBCHEM_BASE}/compound/cid/{cid}/annotations/JSON", timeout=20)
    if not r:
        return None
    try:
        data = r.json()
        full_text = json.dumps(data)
        return clean_cas(full_text)
    except Exception:
        return None


def fetch_from_pubchem(name, smiles_hint, inchikey_hint=None):
    result = {"smiles": None, "cas": None, "iupac": None,
              "formula": None, "mw": None, "inchikey": None,
              "source": None, "cid": None}

    cid = None
    if inchikey_hint:
        r = safe_get(f"{PUBCHEM_BASE}/compound/inchikey/{requests.utils.quote(inchikey_hint, safe='')}/cids/JSON")
        if r and r.status_code == 200:
            try:
                cid = r.json()["IdentifierList"]["CID"][0]
            except Exception:
                pass

    if not cid:
        query_type = "smiles" if smiles_hint else "name"
        query_val = smiles_hint if smiles_hint else name
        r = safe_get(
            f"{PUBCHEM_BASE}/compound/{query_type}/{requests.utils.quote(query_val, safe='')}/cids/JSON"
        )
        if not r or r.status_code != 200:
            return result
        try:
            cid = r.json()["IdentifierList"]["CID"][0]
        except Exception:
            return result

    result["cid"] = cid
    result["source"] = "PubChem"

    r2 = safe_get(
        f"{PUBCHEM_BASE}/compound/cid/{cid}/property/"
        f"IsomericSMILES,IUPACName,MolecularFormula,MolecularWeight,InChIKey/JSON"
    )
    if r2 and r2.status_code == 200:
        try:
            props = r2.json()["PropertyTable"]["Properties"][0]
            result["smiles"] = props.get("IsomericSMILES")
            result["iupac"] = props.get("IUPACName")
            result["formula"] = props.get("MolecularFormula")
            result["mw"] = str(props.get("MolecularWeight", "")) or None
            result["inchikey"] = props.get("InChIKey")
        except Exception:
            pass

    r3 = safe_get(f"{PUBCHEM_BASE}/compound/cid/{cid}/xrefs/RegistryID/JSON")
    if r3 and r3.status_code == 200:
        try:
            ids = r3.json()["InformationList"]["Information"][0]["RegistryID"]
            for rid in ids:
                if re.match(r'^\d{2,7}-\d{2}-\d$', rid):
                    result["cas"] = rid
                    break
        except Exception:
            pass

    if not result["cas"]:
        r4 = safe_get(f"{PUBCHEM_BASE}/compound/cid/{cid}/synonyms/JSON")
        if r4 and r4.status_code == 200:
            try:
                syns = r4.json()["InformationList"]["Information"][0]["Synonym"]
                for s in syns:
                    if re.match(r'^\d{2,7}-\d{2}-\d$', s):
                        result["cas"] = s
                        break
            except Exception:
                pass

    if not result["cas"]:
        cas_pv = fetch_cas_from_pug_view(cid)
        if cas_pv:
            result["cas"] = cas_pv

    if not result["cas"]:
        cas_dep = fetch_cas_from_depositor_xref(cid)
        if cas_dep:
            result["cas"] = cas_dep

    if not result["cas"]:
        cas_ann = fetch_cas_from_annotations(cid)
        if cas_ann:
            result["cas"] = cas_ann

    return result

# ─────────────────────────────────────────────
# Level 2b: NIST Chemistry WebBook
# ─────────────────────────────────────────────

NIST_BASE = "https://webbook.nist.gov/cgi/cbook.cgi"


def fetch_cas_from_nist(inchikey):
    if not inchikey:
        return None
    r = safe_get(
        NIST_BASE,
        params={"InChI": inchikey, "Units": "SI", "Mask": 10},
        timeout=15,
    )
    if not r:
        return None
    m = re.search(
        r'CAS Registry Number:</strong>\s*(\d{2,7}-\d{2}-\d)</li>',
        r.text, re.IGNORECASE,
    )
    if m and validate_cas(m.group(1)):
        return m.group(1)
    return None


# ─────────────────────────────────────────────
# Level 2c: Wikidata SPARQL 批量查询
# ─────────────────────────────────────────────

WIKIDATA_SPARQL = "https://query.wikidata.org/sparql"


def batch_wikidata_cas(inchikey_to_row):
    if not inchikey_to_row:
        return {}
    values = " ".join(f'"{ik}"' for ik in inchikey_to_row)
    query = (
        "SELECT ?inchikey ?cas WHERE {\n"
        f"  VALUES ?inchikey {{ {values} }}\n"
        "  ?compound wdt:P235 ?inchikey.\n"
        "  ?compound wdt:P231 ?cas.\n"
        "}\n"
    )
    r = safe_get(
        WIKIDATA_SPARQL,
        params={"query": query, "format": "json"},
        timeout=30,
    )
    if not r or r.status_code != 200:
        return {}
    result = {}
    try:
        data = r.json()
        bindings = data.get("results", {}).get("bindings", [])
        for b in bindings:
            ik = b.get("inchikey", {}).get("value", "")
            cas = b.get("cas", {}).get("value", "")
            if ik and cas:
                result[ik] = cas
    except Exception:
        pass
    return result


# ─────────────────────────────────────────────
# Level 2d: EPA CompTox Dashboard（兜底）
# ─────────────────────────────────────────────

COMPTOX_ENDPOINTS = [
    "https://comptox.epa.gov/dashboard-api/ccdapp1/chemical/search/equal/{}",
    "https://comptox.epa.gov/dashboard-api/ccdapp1/chemical/search/keyword/{}",
    "https://comptox.epa.gov/dashboardapi/ccdapp1/search?keyword={}",
]


def fetch_cas_from_comptox(inchikey):
    if not inchikey:
        return None
    for endpoint in COMPTOX_ENDPOINTS:
        url = endpoint.format(inchikey)
        r = safe_get(url, timeout=15)
        if r and r.status_code == 200:
            try:
                data = r.json()
                items = data if isinstance(data, list) else data.get("results", [])
                for item in items:
                    cas = item.get("casrn") or item.get("casNumber") or item.get("cas")
                    if cas:
                        cas_str = str(cas).strip()
                        if re.match(r'^\d{2,7}-\d{2}-\d$', cas_str):
                            return cas_str
            except Exception:
                pass
    return None

# ─────────────────────────────────────────────
# 单条查询主逻辑
# ─────────────────────────────────────────────

def query_compound(raw_name):
    raw_name = raw_name.strip()
    if not raw_name:
        return None

    smiles_input = is_smiles(raw_name)
    smiles_hint = raw_name if smiles_input else None
    en_name = "(SMILES)" if smiles_input else translate_name(raw_name)

    row = {
        "原始输入": raw_name,
        "英文名称": "" if smiles_input else en_name,
        "SMILES": None,
        "CAS号": None,
        "IUPAC名称": None,
        "分子式": None,
        "分子量": None,
        "InChIKey": None,
        "数据来源": None,
        "状态": None,
        "cid": None,
        "结构图片": None,
        "_rdkit_ik": None,
    }

    def fill(src, label):
        if not row["SMILES"] and src.get("smiles"):
            row["SMILES"] = src["smiles"]
            row["数据来源"] = label
        if not row["CAS号"] and src.get("cas"):
            row["CAS号"] = src["cas"]
            if label not in (row["数据来源"] or ""):
                row["数据来源"] = (row["数据来源"] or "") + f"+{label}"
        if not row["IUPAC名称"] and src.get("iupac"):
            row["IUPAC名称"] = src["iupac"]
        if not row["分子式"] and src.get("formula"):
            row["分子式"] = src["formula"]
        if not row["分子量"] and src.get("mw"):
            row["分子量"] = src["mw"]
        if not row["InChIKey"] and src.get("inchikey"):
            row["InChIKey"] = src["inchikey"]
        if not row.get("cid") and src.get("cid"):
            row["cid"] = src["cid"]

    # Level 0: RDKit 本地计算 InChIKey
    if smiles_input and _RDKIT_AVAILABLE:
        ik_local, _ = compute_inchikey_local(raw_name)
        if ik_local:
            row["_rdkit_ik"] = ik_local
            if not row["InChIKey"]:
                row["InChIKey"] = ik_local

    # Level 1: CACTUS (SMILES + InChIKey)
    ik_hint = row.get("_rdkit_ik")
    print(f"  [L1-CACTUS] {raw_name[:50]} (IK={'Y' if ik_hint else 'N'})")
    c1 = fetch_from_cactus(raw_name, en_name, smiles_hint, ik_hint)
    fill(c1, c1.get("source") or "CACTUS")

    # Level 2: PubChem
    need_l2 = not row["SMILES"] or not row["CAS号"] or not row["分子式"]
    if need_l2:
        ik_hint = row.get("_rdkit_ik") or row.get("InChIKey")
        print(f"  [L2-PubChem] {raw_name[:50]} (IK={'Y' if ik_hint else 'N'})")
        c2 = fetch_from_pubchem(en_name, row["SMILES"] or smiles_hint, ik_hint)
        fill(c2, "PubChem")

    # 生成结构图片 URL
    if row.get("cid"):
        row["结构图片"] = f"https://pubchem.ncbi.nlm.nih.gov/rest/pug/compound/cid/{row['cid']}/PNG"
    elif row.get("SMILES"):
        row["结构图片"] = f"https://cactus.nci.nih.gov/chemical/structure/{quote(row['SMILES'], safe='')}/image?format=png"

    # CAS 预计算（富集前初判，富集后在 run() 中更新）
    _cf = ["SMILES", "CAS号", "IUPAC名称"]
    _n = sum(1 for k in _cf if row.get(k))
    row["状态"] = "完整" if _n == 3 else ("部分" if _n > 0 and row.get("CAS号") else ("部分(CAS未注册)" if _n > 0 else "失败"))

    return row


# ─────────────────────────────────────────────
# 批量运行 + 进度回调 + 后置 CAS 富集
# ─────────────────────────────────────────────

def run(compounds, progress_callback=None):
    """
    批量查询化合物。
    progress_callback(current, total, compound_name, row_dict, phase)
      phase: 'querying' | 'enriching' | 'done' | 'complete'
    """
    results = []
    total = len(compounds)
    ndjson_path = Path(OUT_DIR) / "results.ndjson"

    for i, name in enumerate(compounds, 1):
        if progress_callback:
            progress_callback(i, total, name, None, 'querying')

        print(f"\n[{i}/{total}] {name[:60]}")
        row = query_compound(name)
        if row:
            results.append(row)
            # 即时落盘，防止中途失败丢失结果
            with open(ndjson_path, 'a', encoding='utf-8') as f:
                f.write(json.dumps(row, ensure_ascii=False) + '\n')

        if progress_callback:
            progress_callback(i, total, name, row, 'done')
        time.sleep(0.3)

    # ── 后置批量 CAS 富集 ──

    # NIST
    missing_cas_nist = [(idx, r) for idx, r in enumerate(results)
                        if not r["CAS号"] and r["InChIKey"]]
    if missing_cas_nist:
        if progress_callback:
            progress_callback(0, total, f"NIST 富集 {len(missing_cas_nist)} 条", None, 'enriching')
        print(f"\n[批量富集] NIST WebBook 查询 {len(missing_cas_nist)} 条...")
        for idx, r in missing_cas_nist:
            ik = r["InChIKey"]
            cas = fetch_cas_from_nist(ik)
            if cas:
                r["CAS号"] = cas
                r["数据来源"] = (r["数据来源"] or "") + "+NIST"
                print(f"  [NIST] {ik} -> CAS {cas}")
            time.sleep(0.2)

    # Wikidata SPARQL 批量查询
    missing_cas_wd = [(idx, r) for idx, r in enumerate(results)
                      if not r["CAS号"] and r["InChIKey"]]
    if missing_cas_wd:
        if progress_callback:
            progress_callback(0, total, f"Wikidata 富集 {len(missing_cas_wd)} 条", None, 'enriching')
        print(f"\n[批量富集] Wikidata SPARQL 查询 {len(missing_cas_wd)} 条...")
        ik_map = {r["InChIKey"]: idx for idx, r in missing_cas_wd}
        wd_results = batch_wikidata_cas(ik_map)
        for ik, cas in wd_results.items():
            idx = ik_map.get(ik)
            if idx is not None and cas:
                results[idx]["CAS号"] = cas
                results[idx]["数据来源"] = (results[idx]["数据来源"] or "") + "+Wikidata"
                print(f"  [Wikidata] {ik} -> CAS {cas}")

    # EPA CompTox
    missing_cas_ct = [(idx, r) for idx, r in enumerate(results)
                      if not r["CAS号"] and r["InChIKey"]]
    if missing_cas_ct:
        if progress_callback:
            progress_callback(0, total, f"CompTox 富集 {len(missing_cas_ct)} 条", None, 'enriching')
        print(f"\n[批量富集] EPA CompTox 查询 {len(missing_cas_ct)} 条...")
        for idx, r in missing_cas_ct:
            ik = r["InChIKey"]
            cas = fetch_cas_from_comptox(ik)
            if cas:
                r["CAS号"] = cas
                r["数据来源"] = (r["数据来源"] or "") + "+CompTox"
                print(f"  [CompTox] {ik} -> CAS {cas}")
            time.sleep(0.3)

    # ── 状态判定 ──
    for r in results:
        core_fields = ["SMILES", "CAS号", "IUPAC名称"]
        filled = sum(1 for k in core_fields if r[k])
        if filled == 3:
            r["状态"] = "完整"
        elif filled > 0:
            if not r["CAS号"]:
                r["状态"] = "部分(CAS未注册)"
            else:
                r["状态"] = "部分"
        else:
            r["状态"] = "失败"
        r.pop("_rdkit_ik", None)
        if r["数据来源"]:
            r["数据来源"] = r["数据来源"].lstrip("+")

    if progress_callback:
        progress_callback(total, total, "全部完成", None, 'complete')

    return results

# ─────────────────────────────────────────────
# CSV / JSON 导出（无 pandas 依赖）
# ─────────────────────────────────────────────

COLUMNS = ["原始输入", "英文名称", "SMILES", "CAS号", "IUPAC名称",
           "分子式", "分子量", "InChIKey", "数据来源", "状态", "结构图片"]


def export_csv(results, filepath=None):
    """导出 CSV，filepath=None 时返回字符串。自动忽略 COLUMNS 以外的字段。"""
    if filepath:
        with open(filepath, 'w', newline='', encoding='utf-8-sig') as f:
            writer = csv.DictWriter(f, fieldnames=COLUMNS, extrasaction='ignore')
            writer.writeheader()
            writer.writerows(results)
        return filepath
    else:
        buf = io.StringIO()
        writer = csv.DictWriter(buf, fieldnames=COLUMNS, extrasaction='ignore')
        writer.writeheader()
        writer.writerows(results)
        return buf.getvalue()


def export_json(results, filepath=None):
    """导出 JSON，filepath=None 时返回字符串"""
    json_str = json.dumps(results, ensure_ascii=False, indent=2)
    if filepath:
        with open(filepath, 'w', encoding='utf-8') as f:
            f.write(json_str)
        return filepath
    return json_str


def parse_input(text):
    """解析用户输入的化合物列表（逗号或换行分隔）"""
    return [c.strip() for c in re.split(r'[,\n]+', text) if c.strip()]


def parse_excel(filepath):
    """从 Excel 文件解析化合物列表，自动检测化合物名称所在列。
    策略：
      1. 扫描第一行，找包含"化合物/name/compound"等关键词的列
      2. 找不到表头关键词，选第一个有 ≥2 个非空非纯数字数据的列
    无需固定格式，数据放任意列均可。
    """
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RuntimeError("需要 openpyxl 库来解析 Excel 文件，请运行: pip install openpyxl")

    wb = load_workbook(filepath, read_only=True, data_only=True)
    ws = wb.active

    all_rows = list(ws.iter_rows(values_only=True))
    if not all_rows:
        wb.close()
        return []

    max_cols = max(len(r) for r in all_rows) if all_rows else 0
    header_keywords = {"化合物名称", "compound", "name", "化合物", "名称",
                       "化合物名", "smiles", "名称列表", "compound name",
                       "chemical", "chemical name", "化合物列表", "drug", "药物名称"}

    # Step 1: 扫描第一行找表头关键词
    target_col = None
    has_header = False
    first_row = all_rows[0] if all_rows else []
    for col_idx in range(min(max_cols, len(first_row))):
        val = first_row[col_idx]
        if val is not None:
            val_str = str(val).strip().lower()
            if val_str in header_keywords:
                target_col = col_idx
                has_header = True
                break

    # Step 2: 找不到表头，选第一个有 ≥2 个非空非纯数字数据的列
    if target_col is None:
        for col_idx in range(max_cols):
            col_values = []
            for row in all_rows:
                if col_idx < len(row) and row[col_idx] is not None:
                    val = str(row[col_idx]).strip()
                    if val:
                        col_values.append(val)
            non_numeric = [v for v in col_values if not re.match(r'^\d+\.?$', v)]
            if len(non_numeric) >= 2:
                target_col = col_idx
                break

    if target_col is None:
        wb.close()
        return []

    # 提取目标列数据，仅当 Step 1 匹配到表头关键词时才跳过第一行
    compounds = []
    start_row = 1 if has_header else 0
    for row in all_rows[start_row:]:
        if target_col < len(row) and row[target_col] is not None:
            val = str(row[target_col]).strip()
            if val:
                compounds.append(val)

    wb.close()
    return compounds


# ─────────────────────────────────────────────
# 入口
# ─────────────────────────────────────────────

if __name__ == "__main__":
    input_str = os.environ.get("COMPOUNDS_INPUT", "")
    if not input_str and len(sys.argv) > 1:
        input_str = " ".join(sys.argv[1:])
    if not input_str:
        print("[ERROR] 请通过 COMPOUNDS_INPUT 环境变量传入化合物列表（换行或逗号分隔）")
        sys.exit(1)

    compounds = parse_input(input_str)
    print(f"共 {len(compounds)} 条化合物")
    if _RDKIT_AVAILABLE:
        print("[RDKit] 已加载，Level 0 InChIKey 本地计算可用")
    else:
        print("[RDKit] 未安装，Level 0 不可用（建议 pip install rdkit）")

    results = run(compounds)

    if not results:
        print("[WARN] 无有效结果")
    else:
        print("\n" + "=" * 80)
        for r in results:
            print(f"  {r.get('英文名称','')[:30]:30s} | SMILES={r.get('SMILES','')[:40]} | CAS={r.get('CAS号','')} | {r.get('状态','')}")
        print("=" * 80)

        out_csv = Path(OUT_DIR) / "compounds_result.csv"
        export_csv(results, str(out_csv))
        print(f"\nCSV 已保存: {out_csv}")

        out_json = Path(OUT_DIR) / "compounds_result.json"
        export_json(results, str(out_json))
        print(f"JSON 已保存: {out_json}")
