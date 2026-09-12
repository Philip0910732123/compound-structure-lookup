#!/usr/bin/env python3
"""API smoke test for running exe"""
import requests, json, os

BASE = "http://127.0.0.1:5173"

# Test 1: Excel upload (data in C column with header)
from openpyxl import Workbook
test_file = "smoke_test.xlsx"
wb = Workbook()
ws = wb.active
ws['A1'] = '序号'; ws['A2'] = 1; ws['A3'] = 2
ws['B1'] = '备注'; ws['B2'] = '无关'
ws['C1'] = '化合物名称'; ws['C2'] = 'aspirin'; ws['C3'] = 'caffeine'
wb.save(test_file)

with open(test_file, "rb") as f:
    r = requests.post(BASE + "/api/upload", files={"file": f}, timeout=10)
os.remove(test_file)

print(f"Upload status: {r.status_code}")
data = r.json()
compounds = data.get("compounds", [])
print(f"Parsed {len(compounds)} compounds: {compounds}")
assert len(compounds) == 2, f"Expected 2, got {len(compounds)}"
assert compounds[0] == 'aspirin'
assert compounds[1] == 'caffeine'
print("[PASS] Excel upload - C column with header detected correctly")

# Test 2: Query with image URL
r2 = requests.post(BASE + "/api/query", json={"compounds": "aspirin"}, stream=True, timeout=60)
assert r2.status_code == 200
buffer = ""
found = None
for chunk in r2.iter_content(chunk_size=1024, decode_unicode=True):
    buffer += chunk
    while "\n" in buffer:
        line, buffer = buffer.split("\n", 1)
        line = line.strip()
        if line.startswith("data: "):
            d = json.loads(line[6:])
            if d.get("phase") == "done" and d.get("row"):
                found = d["row"]
        elif line.startswith("event: done"):
            break
assert found is not None
print(f"Query result - SMILES: {found.get('SMILES')}")
print(f"Query result - CAS: {found.get('CAS号')}")
print(f"Query result - Image: {found.get('结构图片')}")
assert found.get("结构图片"), "Image URL missing"
print("[PASS] Query returns image URL")
print("\n=== All API smoke tests passed ===")
