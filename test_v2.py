#!/usr/bin/env python3
"""功能测试：验证 Excel 解析 + 结构图片 URL 生成"""
import os
import sys
import json

# 确保能 import compound_engine
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from compound_engine import parse_excel, query_compound, COLUMNS

def test_excel_parsing():
    """测试 Excel 解析：创建测试文件 → 解析 → 验证"""
    from openpyxl import Workbook
    
    test_file = os.path.join(os.path.dirname(__file__), "test_input.xlsx")
    wb = Workbook()
    ws = wb.active
    
    # 测试场景1: 数据在 A 列，有表头
    ws['A1'] = '化合物名称'
    ws['A2'] = 'aspirin'
    ws['A3'] = 'caffeine'
    ws['A4'] = 'ibuprofen'
    
    # B 列放一些无关数据
    ws['B1'] = '备注'
    ws['B2'] = '测试'
    
    wb.save(test_file)
    
    compounds = parse_excel(test_file)
    print(f"[Excel解析] 解析到 {len(compounds)} 个化合物: {compounds}")
    assert len(compounds) == 3, f"期望 3 个，实际 {len(compounds)}"
    assert compounds[0] == 'aspirin'
    assert compounds[1] == 'caffeine'
    assert compounds[2] == 'ibuprofen'
    print("[PASS] test_excel_parsing")
    
    # 测试场景2: 数据在 B 列（A 列为空）
    wb2 = Workbook()
    ws2 = wb2.active
    ws2['B1'] = 'osimertinib'
    ws2['B2'] = 'gefitinib'
    wb2.save(test_file)
    
    compounds2 = parse_excel(test_file)
    print(f"[Excel解析-B列] 解析到 {len(compounds2)} 个化合物: {compounds2}")
    assert len(compounds2) == 2
    print("[PASS] test_excel_parsing_b_column")
    
    os.remove(test_file)

def test_image_url():
    """测试结构图片 URL 生成"""
    # 查询 aspirin，验证结构图片 URL
    row = query_compound("aspirin")
    print(f"[图片URL] aspirin 查询结果:")
    print(f"  SMILES: {row.get('SMILES')}")
    print(f"  CAS: {row.get('CAS号')}")
    print(f"  cid: {row.get('cid')}")
    print(f"  结构图片: {row.get('结构图片')}")
    
    assert row.get('结构图片'), "结构图片 URL 不应为空"
    assert 'pubchem' in row.get('结构图片', '') or 'cactus' in row.get('结构图片', ''), \
        f"URL 应来自 PubChem 或 CACTUS，实际: {row.get('结构图片')}"
    print("[PASS] test_image_url")

def test_columns():
    """验证 COLUMNS 包含结构图片"""
    print(f"[COLUMNS] {COLUMNS}")
    assert '结构图片' in COLUMNS, "COLUMNS 应包含 '结构图片'"
    print("[PASS] test_columns")

if __name__ == '__main__':
    print("=== compound-lookup v3.2 功能测试 ===")
    print()
    
    test_columns()
    print()
    
    test_excel_parsing()
    print()
    
    test_image_url()
    print()
    
    print("=== 全部测试通过 ===")
